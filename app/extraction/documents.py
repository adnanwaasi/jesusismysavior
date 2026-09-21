from __future__ import annotations

from abc import ABC, abstractmethod
from pathlib import Path
import string
from uuid import UUID
import zipfile

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image, UnidentifiedImageError
import pypdfium2 as pdfium
from pypdf import PdfReader
import pytesseract

from app.extraction.errors import CorruptDocument, ExtractionError, UnsupportedDocumentType
from app.extraction.models import RawDocument, RawPage


class DocumentExtractor(ABC):
    @abstractmethod
    def supports(self, mime_type: str, path: Path) -> bool: ...

    @abstractmethod
    def extract(self, invoice_id: UUID, mime_type: str, path: Path) -> RawDocument: ...


class OcrEngine:
    def __init__(self, languages: str = "eng", max_image_pixels: int = 80_000_000):
        self.languages = languages
        self.max_image_pixels = max_image_pixels

    def extract_image(self, image: Image.Image) -> str:
        if image.width * image.height > self.max_image_pixels:
            raise ExtractionError("Image exceeds OCR pixel limit")
        return pytesseract.image_to_string(image, lang=self.languages).strip()


def text_is_usable(text: str, minimum_meaningful_characters: int, minimum_printable_ratio: float) -> bool:
    meaningful = [char for char in text if not char.isspace()]
    if len(meaningful) < minimum_meaningful_characters:
        return False
    printable = sum(char in string.printable or char.isprintable() for char in meaningful)
    return printable / len(meaningful) >= minimum_printable_ratio


class PdfDocumentExtractor(DocumentExtractor):
    def __init__(self, ocr: OcrEngine, max_pages: int, min_characters: int, min_printable_ratio: float, ocr_dpi: int):
        self.ocr = ocr
        self.max_pages = max_pages
        self.min_characters = min_characters
        self.min_printable_ratio = min_printable_ratio
        self.ocr_dpi = ocr_dpi

    def supports(self, mime_type: str, path: Path) -> bool:
        return mime_type == "application/pdf" and _starts_with(path, b"%PDF")

    def extract(self, invoice_id: UUID, mime_type: str, path: Path) -> RawDocument:
        if not _starts_with(path, b"%PDF"):
            raise CorruptDocument("PDF signature is missing")
        try:
            reader = PdfReader(path, strict=False)
            if len(reader.pages) > self.max_pages:
                raise ExtractionError(f"PDF exceeds maximum of {self.max_pages} pages")
            text_pages = [RawPage(page_number=index, text=(page.extract_text() or "").strip(), extraction_method="pdf_text") for index, page in enumerate(reader.pages, 1)]
            combined = "\n".join(page.text for page in text_pages)
            if text_is_usable(combined, self.min_characters, self.min_printable_ratio):
                return RawDocument(invoice_id=invoice_id, source_type=mime_type, extraction_method="pdf_text", pages=text_pages, metadata={"page_count": len(text_pages), "ocr_used": False})
            return self._ocr_pdf(invoice_id, mime_type, path)
        except ExtractionError:
            raise
        except Exception as exc:
            raise CorruptDocument(f"Unable to read PDF: {exc}") from exc

    def _ocr_pdf(self, invoice_id: UUID, mime_type: str, path: Path) -> RawDocument:
        try:
            document = pdfium.PdfDocument(path)
            if len(document) > self.max_pages:
                raise ExtractionError(f"PDF exceeds maximum of {self.max_pages} pages")
            pages: list[RawPage] = []
            scale = self.ocr_dpi / 72
            for index in range(len(document)):
                image = document[index].render(scale=scale).to_pil()
                pages.append(RawPage(page_number=index + 1, text=self.ocr.extract_image(image), extraction_method="ocr"))
            return RawDocument(invoice_id=invoice_id, source_type=mime_type, extraction_method="ocr", pages=pages, metadata={"page_count": len(pages), "ocr_used": True, "ocr_languages": self.ocr.languages})
        except ExtractionError:
            raise
        except Exception as exc:
            raise CorruptDocument(f"Unable to OCR PDF: {exc}") from exc


class DocxDocumentExtractor(DocumentExtractor):
    def __init__(self, max_sections: int = 10_000):
        self.max_sections = max_sections

    def supports(self, mime_type: str, path: Path) -> bool:
        if mime_type != "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            return False
        try:
            with zipfile.ZipFile(path) as archive:
                return "word/document.xml" in archive.namelist()
        except zipfile.BadZipFile:
            return False

    def extract(self, invoice_id: UUID, mime_type: str, path: Path) -> RawDocument:
        try:
            document = Document(path)
            blocks: list[str] = []
            for child in document.element.body.iterchildren():
                if child.tag.endswith("}p"):
                    text = Paragraph(child, document).text.strip()
                elif child.tag.endswith("}tbl"):
                    table = Table(child, document)
                    text = "\n".join(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
                else:
                    continue
                if text:
                    blocks.append(text)
                if len(blocks) > self.max_sections:
                    raise ExtractionError("DOCX exceeds section limit")
            page = RawPage(page_number=1, text="\n".join(blocks), extraction_method="docx")
            return RawDocument(invoice_id=invoice_id, source_type=mime_type, extraction_method="docx", pages=[page], metadata={"section_count": len(blocks), "page_count": 1})
        except ExtractionError:
            raise
        except Exception as exc:
            raise CorruptDocument(f"Unable to read DOCX: {exc}") from exc


class ImageDocumentExtractor(DocumentExtractor):
    MIME_TYPES = frozenset({"image/png", "image/jpeg", "image/tiff", "image/bmp"})

    def __init__(self, ocr: OcrEngine):
        self.ocr = ocr

    def supports(self, mime_type: str, path: Path) -> bool:
        if mime_type not in self.MIME_TYPES:
            return False
        try:
            with Image.open(path) as image:
                image.verify()
            return True
        except (UnidentifiedImageError, OSError):
            return False

    def extract(self, invoice_id: UUID, mime_type: str, path: Path) -> RawDocument:
        try:
            with Image.open(path) as image:
                text = self.ocr.extract_image(image.convert("RGB"))
            page = RawPage(page_number=1, text=text, extraction_method="ocr")
            return RawDocument(invoice_id=invoice_id, source_type=mime_type, extraction_method="ocr", pages=[page], metadata={"page_count": 1, "ocr_used": True, "ocr_languages": self.ocr.languages})
        except ExtractionError:
            raise
        except Exception as exc:
            raise CorruptDocument(f"Unable to OCR image: {exc}") from exc


class DocumentExtractorDispatcher:
    def __init__(self, extractors: list[DocumentExtractor]):
        self.extractors = extractors

    def select(self, mime_type: str, path: Path) -> DocumentExtractor:
        for extractor in self.extractors:
            if extractor.supports(mime_type, path):
                return extractor
        raise UnsupportedDocumentType(f"Unsupported or invalid document type: {mime_type}")

    def extract(self, invoice_id: UUID, mime_type: str, path: Path) -> RawDocument:
        return self.select(mime_type, path).extract(invoice_id, mime_type, path)


def _starts_with(path: Path, signature: bytes) -> bool:
    try:
        with path.open("rb") as source:
            return source.read(len(signature)) == signature
    except OSError:
        return False
