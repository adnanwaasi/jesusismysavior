from __future__ import annotations

from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw, ImageFont
from reportlab.pdfgen import canvas


INVOICE_LINES = [
    "Invoice Number: INV-1001",
    "Invoice Date: 2026-09-20",
    "Vendor: Example Supplier",
    "Currency: USD",
    "Description | Quantity | Unit Price | Line Total",
    "Item A | 2 | 200.00 | 400.00",
    "Item B | 3 | 200.00 | 600.00",
    "Subtotal: 1000.00",
    "Tax: 100.00",
    "Total: 1100.00",
]


def make_digital_pdf(path: Path) -> Path:
    document = canvas.Canvas(str(path), pagesize=(612, 792))
    y = 740
    for line in INVOICE_LINES:
        document.drawString(60, y, line)
        y -= 28
    document.save()
    return path


def make_invoice_image(path: Path) -> Path:
    image = Image.new("RGB", (1600, 1200), "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 34)
    except OSError:
        font = ImageFont.load_default()
    y = 60
    for line in INVOICE_LINES:
        draw.text((70, y), line, fill="black", font=font)
        y += 90
    image.save(path)
    return path


def make_scanned_pdf(path: Path) -> Path:
    image_path = path.with_suffix(".source.png")
    image = Image.open(make_invoice_image(image_path))
    image.save(path, "PDF", resolution=150.0)
    image_path.unlink()
    return path


def make_docx(path: Path) -> Path:
    document = Document()
    for line in INVOICE_LINES[:4]:
        document.add_paragraph(line)
    table = document.add_table(rows=1, cols=4)
    for cell, value in zip(table.rows[0].cells, ("Description", "Quantity", "Unit Price", "Line Total")):
        cell.text = value
    for row_values in (("Item A", "2", "200.00", "400.00"), ("Item B", "3", "200.00", "600.00")):
        row = table.add_row()
        for cell, value in zip(row.cells, row_values):
            cell.text = value
    for line in INVOICE_LINES[-3:]:
        document.add_paragraph(line)
    document.save(path)
    return path
